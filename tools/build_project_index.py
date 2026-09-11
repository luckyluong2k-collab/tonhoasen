"""Build a browser-loadable, offline search index for the Hoa Sen project files."""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
from pathlib import Path
from urllib.parse import quote


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "assets" / "hoa-sen" / "docs"
OUTPUT = ROOT / "assets" / "hoa-sen" / "project-document-index.js"
MAX_TEXT = 6500


def clean(value: object) -> str:
    text = str(value or "")
    text = text.replace("\x00", " ").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def shorten(text: str, limit: int = MAX_TEXT) -> str:
    text = clean(text)
    return text if len(text) <= limit else text[:limit].rsplit(" ", 1)[0] + " ..."


def url_for(path: Path) -> str:
    relative = path.relative_to(ROOT).as_posix()
    return "./" + "/".join(quote(part, safe="") for part in relative.split("/"))


def record(path: Path, kind: str, locator: str, text: str, **extra: object) -> dict:
    return {
        "source": path.name,
        "url": url_for(path),
        "kind": kind,
        "locator": locator,
        "text": shorten(text),
        **extra,
    }


def add_text_chunks(index: list[dict], path: Path, kind: str, text: str) -> None:
    paragraphs = [clean(item) for item in re.split(r"\n\s*\n|\n", text) if clean(item)]
    if not paragraphs:
        return
    chunk: list[str] = []
    size = 0
    start = 1
    for number, paragraph in enumerate(paragraphs, 1):
        if chunk and size + len(paragraph) > MAX_TEXT:
            index.append(record(path, kind, f"đoạn {start}-{number - 1}", "\n".join(chunk), paragraphStart=start, paragraphEnd=number - 1))
            chunk = []
            size = 0
            start = number
        chunk.append(paragraph)
        size += len(paragraph) + 1
    if chunk:
        index.append(record(path, kind, f"đoạn {start}-{len(paragraphs)}", "\n".join(chunk), paragraphStart=start, paragraphEnd=len(paragraphs)))


def read_xlsx(path: Path, index: list[dict]) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            values = [clean(value) for value in row]
            while values and not values[-1]:
                values.pop()
            if any(values):
                index.append(record(path, "xlsx", f"sheet {sheet.title}, dòng {row_number}", " | ".join(values), sheet=sheet.title, row=row_number))
    workbook.close()


def read_xls(path: Path, index: list[dict]) -> None:
    xlrd_path = os.environ.get("HSH_XLRD_PATH")
    if xlrd_path:
        sys.path.insert(0, xlrd_path)
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    for sheet in workbook.sheets():
        for row_number in range(sheet.nrows):
            values = [clean(value) for value in sheet.row_values(row_number)]
            while values and not values[-1]:
                values.pop()
            if any(values):
                index.append(record(path, "xls", f"sheet {sheet.name}, dòng {row_number + 1}", " | ".join(values), sheet=sheet.name, row=row_number + 1))
    workbook.release_resources()


def read_docx(path: Path, index: list[dict]) -> None:
    from docx import Document

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table_number, table in enumerate(document.tables, 1):
        parts.append(f"[Bảng {table_number}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    add_text_chunks(index, path, "docx", "\n".join(parts))


def read_pdf(path: Path, index: list[dict]) -> None:
    from pypdf import PdfReader

    reader = PdfReader(path)
    extracted = 0
    for page_number, page in enumerate(reader.pages, 1):
        text = clean(page.extract_text() or "")
        if text:
            index.append(record(path, "pdf", f"trang {page_number}", text, page=page_number))
            extracted += 1
    if not extracted:
        index.append(record(path, "pdf-scan", f"{len(reader.pages)} trang ảnh/scan", "PDF không có lớp chữ để tìm kiếm tự động; vẫn liên kết được tài liệu gốc để mở và xem."))


def read_legacy_with_office(path: Path) -> str:
    """Use installed Office only while building the static index; the web app stays offline."""
    escaped = str(path).replace("'", "''")
    if path.suffix.lower() == ".doc":
        script = f"""
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$doc = $word.Documents.Open('{escaped}', $false, $true)
[Console]::Write($doc.Content.Text)
$doc.Close($false)
$word.Quit()
"""
    else:
        script = f"""
$excel = New-Object -ComObject Excel.Application
$excel.Visible = $false
$book = $excel.Workbooks.Open('{escaped}', $false, $true)
foreach ($sheet in $book.Worksheets) {{
  Write-Output ('[Sheet ' + $sheet.Name + ']')
  $range = $sheet.UsedRange
  for ($r = 1; $r -le $range.Rows.Count; $r++) {{
    $values = @()
    for ($c = 1; $c -le $range.Columns.Count; $c++) {{ $values += [string]$range.Cells.Item($r, $c).Text }}
    Write-Output ($values -join ' | ')
  }}
}}
$book.Close($false)
$excel.Quit()
"""
    result = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", script], capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180)
    if result.returncode:
        raise RuntimeError(result.stderr.strip() or "Office extraction failed")
    return result.stdout


def main() -> None:
    index: list[dict] = []
    manifest: list[dict] = []
    readers = {".xlsx": read_xlsx, ".xls": read_xls, ".docx": read_docx, ".pdf": read_pdf}
    for path in sorted(DOCS.iterdir(), key=lambda item: item.name.lower()):
        if not path.is_file():
            continue
        before = len(index)
        try:
            if path.suffix.lower() in readers:
                readers[path.suffix.lower()](path, index)
                status = "indexed"
            elif path.suffix.lower() == ".doc":
                add_text_chunks(index, path, "doc", read_legacy_with_office(path))
                status = "indexed-office"
            else:
                status = "linked-only"
        except Exception as exc:  # Keep the file discoverable even if one parser fails.
            index.append(record(path, "metadata", "toàn bộ tài liệu", f"Không đọc được nội dung tự động: {exc}. Mở tài liệu gốc để kiểm tra."))
            status = f"metadata-only: {type(exc).__name__}"
        manifest.append({"source": path.name, "kind": path.suffix.lower().lstrip("."), "status": status, "records": len(index) - before, "url": url_for(path)})

    payload = "window.PROJECT_DOCUMENT_INDEX = " + json.dumps(index, ensure_ascii=False, separators=(",", ":")) + ";\n"
    payload += "window.PROJECT_DOCUMENT_MANIFEST = " + json.dumps(manifest, ensure_ascii=False, separators=(",", ":")) + ";\n"
    OUTPUT.write_text(payload, encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), "records": len(index), "files": manifest}, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
